"""Regression checks against main.py without launching the Streamlit UI."""
import ast
import re
import unittest
from collections import Counter, defaultdict
from copy import copy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
import xlsxwriter

source = Path(__file__).with_name('main.py').read_text(encoding='utf-8')
compile(source, 'main.py', 'exec')
tree = ast.parse(source)
nodes = []
for node in tree.body:
    if isinstance(node, ast.FunctionDef):
        node.decorator_list = []
        nodes.append(node)
    elif isinstance(node, ast.Assign) and all(
        isinstance(target, ast.Name) and target.id.isupper() for target in node.targets
    ):
        nodes.append(node)
ns = dict(globals(), pd=SimpleNamespace(isna=lambda value: value is None), fitz=None)
exec(compile(ast.Module(body=nodes, type_ignores=[]), 'main.py', 'exec'), ns)


class Document(list):
    def close(self):
        pass


class Page:
    rect = SimpleNamespace(height=700)
    def __init__(self, text, words=()):
        self.text, self.words = text, words
    def get_text(self, kind, **kwargs):
        return self.words if kind == 'words' else self.text


def word(x, y, text):
    return (x, y, x+10, y+8, text, 0, 0, 0)


class ReviewFixes(unittest.TestCase):
    def test_mpn_spacing_in_every_import_path(self):
        note = ('Basic Data Text:\nOPERATING TEMPERATURE: -50 TO 60DEG.C\n'
                'STANDARD: IEC 60623\n\nMPN Text:\n'
                'Manufacturer Part Number: VN 570-1\n'
                'Manufacturer number: 0020001566 - ALCAD LTD.')
        template = Workbook()
        data = ns['workbook_bytes'](template)
        for style in ('default', 'tabular_pdf', 'nawah_pdf', 'nawah_excel', 'dubai_csv'):
            with self.subTest(style=style):
                ns['st'] = SimpleNamespace(session_state={})
                ns['prepare_review']([('10', 'BATTERY', '1', 'EA', note)],
                                     BytesIO(data), 'test', 'test', format_style=style)
                draft = ns['st'].session_state['test_review']['queue'][0]['formatted']
                self.assertIn('MPN Text:\n\nManufacturer Part Number: VN 570-1\n'
                              'Manufacturer number: 0020001566 - ALCAD LTD.', draft)

    def test_mpn_spacing_is_idempotent(self):
        space = ns['space_after_initial_section_header']
        for header in ('MPN Text:', 'MPN Details', '**MPN Text:**'):
            for blank in ('', '\n', '\n\n'):
                text = 'STANDARD: IEC 60623\n\n' + header + '\n' + blank + 'Manufacturer Part Number: VN 570-1\nManufacturer number: 123 - ACME'
                expected = 'STANDARD: IEC 60623\n\n' + header + '\n\nManufacturer Part Number: VN 570-1\nManufacturer number: 123 - ACME'
                self.assertEqual(space(text), expected)
                self.assertEqual(space(space(text)), expected)

    def test_navigation_saves_edits_and_returns_to_preview(self):
        queue = [dict(source='Original', formatted='Original', manufacturer='ACME',
                      desc='PUMP', editing=True, tabular_pdf=True),
                 dict(source='Second', formatted='Second', manufacturer='',
                      desc='VALVE', editing=False)]
        review = dict(queue=queue, position=0)
        state = {'editable_test_0': '**Revised note**',
                 'editablemaker_test_0': 'Updated maker',
                 'desc_test_0': 'Updated description'}
        reruns = []
        ns['st'] = SimpleNamespace(session_state=state, rerun=lambda **kw: reruns.append(kw))
        ns['_navigate_review'](review, 'test', 1)
        self.assertFalse(queue[0]['editing'])
        self.assertEqual(queue[0]['formatted'], '**Revised note**')
        self.assertEqual(queue[0]['manufacturer'], 'Updated maker')
        self.assertEqual(queue[0]['desc'], 'Updated description')
        ns['_navigate_review'](review, 'test', 0)
        self.assertEqual(review['position'], 0)
        self.assertFalse(queue[0]['editing'])
        self.assertEqual(state['formatted_test_0'], '**Revised note**')
        self.assertEqual(len(reruns), 2)
        state['test_review'] = review
        ns['_review_enter_edit']('test')
        self.assertTrue(queue[0]['editing'])
        self.assertEqual(state['editable_test_0'], '**Revised note**')

    def test_navigation_with_empty_editor_stays_on_current_item(self):
        item = dict(source='Original', formatted='Original', manufacturer='', editing=True)
        review = dict(queue=[item, dict(editing=False)], position=0)
        errors = []
        ns['st'] = SimpleNamespace(session_state={'editable_test_0': ''}, error=errors.append)
        ns['_navigate_review'](review, 'test', 1)
        self.assertEqual(review['position'], 0)
        self.assertTrue(item['editing'])
        self.assertEqual(item['formatted'], 'Original')
        self.assertTrue(errors)

    def test_technical_punctuation_and_boundaries(self):
        contains = ns['_multiset_contains']
        for expected, changed in [('1.5', '15'), ('AB-12', 'AB12'),
                                  ('123', '91234'), ('1/2', '12'),
                                  ('-5', '5'), ('+5', '5'), ('ABC', 'XABC')]:
            with self.subTest(expected=expected):
                self.assertFalse(contains(expected, changed))
        self.assertTrue(contains('AB-12', 'PART NUMBER: AB-12\n'))
        self.assertTrue(contains('PRESSURE : 1.5 BAR', 'Pressure: 1.5 BAR'))
        self.assertTrue(contains('ACME LTD.', 'MAKE: ACME LTD.'))

    def test_validator_flags_changes(self):
        for source, draft in [('Manufacturer Part Number: AB-12', 'Manufacturer Part Number: AB12'),
                              ('Pressure: 1.5 BAR', 'Pressure: 15 BAR'),
                              ('Operating limit -5 BAR', 'Operating limit 5 BAR')]:
            with self.subTest(source=source):
                self.assertTrue(ns['validate_rfq'](source, draft))

    def test_nawah_quality_preserved_and_checked(self):
        note = ('Material PO Text:\nPUMP\nQUALITY REQUIREMENTS\n'
                'START OF MESSAGE CODE Z_NPE_UNLISTED\n'
                'Supplier must submit document XYZ-9 within 15 days.\n'
                'PACKAGING: moisture barrier required\nEND OF MESSAGE CODE')
        for style in ('nawah_pdf', 'nawah_excel'):
            draft, _ = ns['summarize_nawah_working_note'](note)
            self.assertEqual(draft, note)
            self.assertEqual(ns['validate_rfq'](note, draft, style), [])
            changed = draft.replace('Supplier must submit document XYZ-9 within 15 days.\n', '')
            self.assertTrue(any(w['priority'] == 'Critical'
                                for w in ns['validate_rfq'](note, changed, style)))
            self.assertEqual(ns['_classification']('PACKAGING: moisture barrier required', changed, style), 'matched')
            self.assertEqual(ns['_classification']('PACKAGING: moisture barrier required', 'PUMP', style), 'missing')

    def adnoc(self, item='00010', short='Short Text: PUMP', quantity='2'):
        return f'{item}\n120000000001\n{quantity}\nEA\n0\n0\n0\n23.10.2026\n{short}\nPO Material Text:\nManufacturer Part Number: AB-12\nAgreement / LineNo.\n'

    def test_adnoc_complete(self):
        text = self.adnoc() + self.adnoc('00020')
        self.assertEqual(len(ns['parse_pdf'](text, text)), 2)

    def test_adnoc_partial_rejected(self):
        for bad in (self.adnoc('00020', short='No description'),
                    self.adnoc('00020', quantity='unknown')):
            text = self.adnoc() + bad
            with self.subTest(bad=bad), self.assertRaisesRegex(ValueError, 'partial import'):
                ns['parse_pdf'](text, text)

    def test_adnoc_legacy_missing_note_rejected(self):
        text = ('00010 120000000001 2 EA 0 23.10.2026\nShort Text: PUMP\n'
                'PO Material Text: AB-12\nAgreement / LineNo.\n'
                '00020 120000000002 2 EA 0 23.10.2026\nShort Text: VALVE\n')
        with self.assertRaisesRegex(ValueError, '00020'):
            ns['parse_pdf'](text, text)

    def test_adnoc_mixed_layout_cannot_drop_inline_item(self):
        text = self.adnoc() + '00020 120000000002 2 EA 0 23.10.2026\nShort Text: VALVE\nPO Material Text: NOTE\nAgreement / LineNo.\n'
        with self.assertRaisesRegex(ValueError, 'counts differ'):
            ns['parse_pdf'](text, text)

    def test_fertil_complete_and_rejected_rows(self):
        words = [word(200, 100, 'Description')]
        for y, num in ((150, '1'), (250, '2')):
            words += [word(10, y, num), word(90, y, '10000'),
                      word(120, y, '123456789'), word(200, y, 'PUMP'),
                      word(410, y, '2'), word(450, y, 'EA')]
        def parse(words):
            ns['fitz'] = SimpleNamespace(open=lambda **kw: Document([Page('RFQ FORM 12345678', words)]))
            return ns['parse_tabular_rfq_pdf'](BytesIO(b'PDF'))
        self.assertEqual(len(parse(words)), 2)
        missing_qty = [w for w in words if not (w[0] == 410 and w[1] == 250)]
        with self.assertRaisesRegex(ValueError, 'Fertil item 2'):
            parse(missing_qty)

    def test_nawah_partial_row_rejected(self):
        text = ('NAWAH ENERGY\nMPN NO\n1 123456789 PUMP  AB-12  2 EA 23 Oct 2026\n'
                '[Material PO Text]\nPUMP NOTE\n')
        def parse(text):
            ns['fitz'] = SimpleNamespace(open=lambda **kw: Document([Page(text)]))
            return ns['parse_nawah_rfq_pdf'](BytesIO(b'PDF'))
        self.assertEqual(len(parse(text)), 1)
        with self.assertRaisesRegex(ValueError, 'partial import'):
            parse(text + '2 123456790 VALVE AB-13 BAD EA 23 Oct 2026\n')

    def test_excel_text_limit(self):
        check = ns['checked_excel_value']
        self.assertEqual(len(check('x'*32767)), 32767)
        with self.assertRaisesRegex(ValueError, '32,767'):
            check('x'*32768)
        with self.assertRaises(ValueError):
            check('\U0001f600'*16384)

    def test_rich_export_limit_and_return_code(self):
        item = {'line': '10', 'formatted': 'x'*32768}
        with self.assertRaisesRegex(ValueError, '32,767'):
            ns['make_rich_text_shared_strings'](b'', [item])
        item['formatted'] = 'valid text'
        with patch.object(xlsxwriter.worksheet.Worksheet, 'write_string', return_value=-2):
            with self.assertRaisesRegex(ValueError, 'code -2'):
                ns['make_rich_text_shared_strings'](b'', [item])

    def test_export_round_trip(self):
        template = Workbook()
        template.active.append(['Line', 'Description', 'Qty', 'UOM', 'Note'])
        queue = [dict(line='10', desc='PUMP', qty='2', uom='EA', verified=True,
                      formatted='MPN Text:\n\nManufacturer Part Number: AB-12\nManufacturer number: 002 - ACME'),
                 dict(line='20', desc='VALVE', qty=1.0, uom='EA', verified=True,
                      formatted='x'*32767)]
        review = dict(queue=queue, template=ns['workbook_bytes'](template), upload_bytes=None, suffix='test')
        ns['st'] = SimpleNamespace(session_state={'test_review': review})
        ns['finalize_verified']('test')
        output = ns['st'].session_state['test_final_bytes']
        wb = load_workbook(BytesIO(output), rich_text=True)
        self.assertEqual(str(wb.active['E2'].value), queue[0]['formatted'])
        self.assertTrue(any(isinstance(run, TextBlock) and run.font.b for run in wb.active['E2'].value))
        self.assertEqual(str(wb.active['E3'].value), queue[1]['formatted'])
        wb.close()

    def test_upload_blank_and_multiline_round_trip(self):
        wb = Workbook()
        wb.active.append(['RFQ', 'Note', 'Optional'])
        wb.active.append(['12345678', 'Line one\nLine two', ''])
        data = ns['workbook_bytes'](wb)
        saved = load_workbook(BytesIO(data))
        self.assertEqual(saved.active['B2'].value, 'Line one\nLine two')
        self.assertIsNone(saved.active['C2'].value)
        saved.close()

    def test_export_content_mismatch_blocks_delivery(self):
        wb = Workbook()
        wb.active.append(['Line', 'Description', 'Qty', 'UOM', 'Note'])
        item = dict(line='10', desc='PUMP', qty='2', uom='EA', verified=True, formatted='Required text')
        review = dict(queue=[item], template=ns['workbook_bytes'](wb), upload_bytes=None, suffix='test')
        ns['st'] = SimpleNamespace(session_state={'test_review': review})
        def corrupt(data, items):
            saved = load_workbook(BytesIO(data))
            saved.active['E2'] = 'Changed text'
            return ns['workbook_bytes'](saved)
        original = ns['make_rich_text_shared_strings']
        try:
            ns['make_rich_text_shared_strings'] = corrupt
            with self.assertRaisesRegex(ValueError, 'Export changed item'):
                ns['finalize_verified']('test')
            self.assertNotIn('test_final_bytes', ns['st'].session_state)
        finally:
            ns['make_rich_text_shared_strings'] = original


if __name__ == '__main__':
    unittest.main(verbosity=2)
